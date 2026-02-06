"""Document loader supporting multiple formats with metadata extraction."""

from __future__ import annotations

import hashlib
import mimetypes
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Any, BinaryIO

import structlog

logger = structlog.get_logger(__name__)


class DocumentFormat(str, Enum):
    PDF = "pdf"
    DOCX = "docx"
    TXT = "txt"
    MARKDOWN = "markdown"
    HTML = "html"


@dataclass
class DocumentMetadata:
    """Metadata extracted from a document."""

    source: str
    format: DocumentFormat
    title: str = ""
    author: str = ""
    created_at: datetime | None = None
    page_count: int | None = None
    word_count: int = 0
    content_hash: str = ""
    extra: dict[str, Any] = field(default_factory=dict)


@dataclass
class LoadedDocument:
    """A document loaded into memory with content and metadata."""

    id: str
    content: str
    metadata: DocumentMetadata
    pages: list[str] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return len(self.content.strip()) == 0


def _compute_hash(content: str) -> str:
    return hashlib.sha256(content.encode("utf-8")).hexdigest()[:16]


def _generate_doc_id(source: str, content_hash: str) -> str:
    return f"doc_{hashlib.md5(f'{source}:{content_hash}'.encode()).hexdigest()[:12]}"


class PDFLoader:
    """Load PDF documents using pypdf."""

    @staticmethod
    def load(source: str | Path, stream: BinaryIO | None = None) -> LoadedDocument:
        from pypdf import PdfReader

        if stream is not None:
            reader = PdfReader(stream)
        else:
            reader = PdfReader(str(source))

        pages: list[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            pages.append(text)

        content = "\n\n".join(pages)
        content_hash = _compute_hash(content)

        pdf_meta = reader.metadata or {}
        title = str(pdf_meta.get("/Title", "")) if pdf_meta.get("/Title") else ""
        author = str(pdf_meta.get("/Author", "")) if pdf_meta.get("/Author") else ""
        created_str = str(pdf_meta.get("/CreationDate", ""))
        created_at = _parse_pdf_date(created_str) if created_str else None

        metadata = DocumentMetadata(
            source=str(source),
            format=DocumentFormat.PDF,
            title=title,
            author=author,
            created_at=created_at,
            page_count=len(reader.pages),
            word_count=len(content.split()),
            content_hash=content_hash,
        )

        return LoadedDocument(
            id=_generate_doc_id(str(source), content_hash),
            content=content,
            metadata=metadata,
            pages=pages,
        )


class DOCXLoader:
    """Load DOCX documents using python-docx."""

    @staticmethod
    def load(source: str | Path, stream: BinaryIO | None = None) -> LoadedDocument:
        from docx import Document as DocxDocument

        if stream is not None:
            doc = DocxDocument(stream)
        else:
            doc = DocxDocument(str(source))

        paragraphs: list[str] = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)

        content = "\n\n".join(paragraphs)
        content_hash = _compute_hash(content)

        core = doc.core_properties
        metadata = DocumentMetadata(
            source=str(source),
            format=DocumentFormat.DOCX,
            title=core.title or "",
            author=core.author or "",
            created_at=core.created,
            word_count=len(content.split()),
            content_hash=content_hash,
        )

        return LoadedDocument(
            id=_generate_doc_id(str(source), content_hash),
            content=content,
            metadata=metadata,
        )


class TextLoader:
    """Load plain text files."""

    @staticmethod
    def load(
        source: str | Path,
        stream: BinaryIO | None = None,
        fmt: DocumentFormat = DocumentFormat.TXT,
    ) -> LoadedDocument:
        if stream is not None:
            content = stream.read().decode("utf-8", errors="replace")
        else:
            content = Path(source).read_text(encoding="utf-8", errors="replace")

        content_hash = _compute_hash(content)

        title = ""
        if fmt == DocumentFormat.MARKDOWN:
            match = re.search(r"^#\s+(.+)$", content, re.MULTILINE)
            if match:
                title = match.group(1).strip()
        elif fmt == DocumentFormat.HTML:
            match = re.search(r"<title>(.*?)</title>", content, re.IGNORECASE | re.DOTALL)
            if match:
                title = match.group(1).strip()

        metadata = DocumentMetadata(
            source=str(source),
            format=fmt,
            title=title,
            word_count=len(content.split()),
            content_hash=content_hash,
        )

        return LoadedDocument(
            id=_generate_doc_id(str(source), content_hash),
            content=content,
            metadata=metadata,
        )


# --- Format detection ---

_EXTENSION_MAP: dict[str, DocumentFormat] = {
    ".pdf": DocumentFormat.PDF,
    ".docx": DocumentFormat.DOCX,
    ".txt": DocumentFormat.TXT,
    ".md": DocumentFormat.MARKDOWN,
    ".markdown": DocumentFormat.MARKDOWN,
    ".html": DocumentFormat.HTML,
    ".htm": DocumentFormat.HTML,
}

_MIME_MAP: dict[str, DocumentFormat] = {
    "application/pdf": DocumentFormat.PDF,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": DocumentFormat.DOCX,
    "text/plain": DocumentFormat.TXT,
    "text/markdown": DocumentFormat.MARKDOWN,
    "text/html": DocumentFormat.HTML,
}


def detect_format(
    source: str | Path,
    mime_type: str | None = None,
) -> DocumentFormat:
    """Detect document format from file extension or MIME type."""
    if mime_type:
        fmt = _MIME_MAP.get(mime_type)
        if fmt:
            return fmt

    ext = Path(source).suffix.lower()
    fmt = _EXTENSION_MAP.get(ext)
    if fmt:
        return fmt

    guessed, _ = mimetypes.guess_type(str(source))
    if guessed:
        fmt = _MIME_MAP.get(guessed)
        if fmt:
            return fmt

    return DocumentFormat.TXT


def _parse_pdf_date(date_str: str) -> datetime | None:
    """Parse PDF date strings like D:20230101120000+00'00'."""
    if not date_str:
        return None
    try:
        cleaned = date_str.replace("D:", "").replace("'", "")
        if len(cleaned) >= 14:
            return datetime.strptime(cleaned[:14], "%Y%m%d%H%M%S").replace(
                tzinfo=timezone.utc
            )
        elif len(cleaned) >= 8:
            return datetime.strptime(cleaned[:8], "%Y%m%d").replace(
                tzinfo=timezone.utc
            )
    except (ValueError, IndexError):
        pass
    return None


class DocumentLoader:
    """Unified document loader with automatic format detection.

    Supports PDF, DOCX, TXT, Markdown, and HTML documents.
    Extracts metadata including title, author, date, and word count.
    """

    _loaders: dict[DocumentFormat, type] = {
        DocumentFormat.PDF: PDFLoader,
        DocumentFormat.DOCX: DOCXLoader,
        DocumentFormat.TXT: TextLoader,
        DocumentFormat.MARKDOWN: TextLoader,
        DocumentFormat.HTML: TextLoader,
    }

    def load(
        self,
        source: str | Path,
        *,
        mime_type: str | None = None,
        stream: BinaryIO | None = None,
    ) -> LoadedDocument:
        """Load a document from a file path or byte stream.

        Args:
            source: File path or identifier for the document.
            mime_type: Optional MIME type for explicit format specification.
            stream: Optional binary stream (for uploaded files).

        Returns:
            A LoadedDocument with content, pages, and extracted metadata.

        Raises:
            ValueError: If the document format is unsupported.
            FileNotFoundError: If the source path does not exist.
        """
        fmt = detect_format(source, mime_type)
        logger.info("loading_document", source=str(source), format=fmt.value)

        if stream is None and not Path(source).exists():
            raise FileNotFoundError(f"Document not found: {source}")

        loader_cls = self._loaders.get(fmt)
        if loader_cls is None:
            raise ValueError(f"Unsupported format: {fmt}")

        if loader_cls is TextLoader:
            return TextLoader.load(source, stream=stream, fmt=fmt)
        return loader_cls.load(source, stream=stream)

    def load_directory(
        self,
        directory: str | Path,
        *,
        recursive: bool = True,
        extensions: set[str] | None = None,
    ) -> list[LoadedDocument]:
        """Load all supported documents from a directory.

        Args:
            directory: Path to the directory to scan.
            recursive: Whether to recurse into subdirectories.
            extensions: Optional set of extensions to filter by (e.g., {".pdf", ".md"}).

        Returns:
            List of loaded documents.
        """
        root = Path(directory)
        if not root.is_dir():
            raise NotADirectoryError(f"Not a directory: {directory}")

        pattern = "**/*" if recursive else "*"
        documents: list[LoadedDocument] = []

        for path in sorted(root.glob(pattern)):
            if not path.is_file():
                continue
            if extensions and path.suffix.lower() not in extensions:
                continue
            if path.suffix.lower() not in _EXTENSION_MAP:
                continue

            try:
                doc = self.load(path)
                if not doc.is_empty:
                    documents.append(doc)
            except Exception:
                logger.warning("failed_to_load", path=str(path), exc_info=True)

        logger.info("directory_loaded", directory=str(directory), count=len(documents))
        return documents
